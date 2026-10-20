"""Texto de los documentos de la biblioteca, para el buscador."""

import io
import logging
import re

import docx
from pypdf import PdfReader

logger = logging.getLogger(__name__)

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
# Muy por encima de lo que da un documento de 10 MB; evita pasar del límite del índice
MAX_CHARS = 1_000_000


def _pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(content: bytes) -> str:
    document = docx.Document(io.BytesIO(content))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(lines)


def extract_text(content: bytes, mime_type: str) -> str:
    """Texto del PDF o .docx; vacío si no tiene (PDF escaneado) o no se puede leer."""
    try:
        text = _pdf_text(content) if mime_type == PDF else _docx_text(content)
    except Exception:  # un fichero dañado no impide guardarlo
        logger.warning("No se pudo extraer el texto de un %s", mime_type, exc_info=True)
        return ""
    # Postgres no admite el carácter nulo en el texto
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s*\n\s*", "\n", text).strip()
    return text[:MAX_CHARS]
