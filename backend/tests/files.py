"""Ficheros de prueba con texto para la biblioteca."""

import io

import docx


def pdf_with_text(*lines: str) -> bytes:
    """PDF mínimo de una página con esas líneas (fuente estándar, admite tildes)."""
    stream = b"BT /F1 12 Tf 72 720 Td 14 TL "
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream += b"(" + escaped.encode("cp1252") + b") ' "
    stream += b"ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return out


def docx_with_text(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        cells = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(cells.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
